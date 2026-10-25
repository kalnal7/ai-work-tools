import csv
import io
import json
from openpyxl import Workbook, load_workbook
import pytest
import app


LABELS = list(app.TOPICS)


def test_fixed_30_review_expected_types_compounds_negation_and_unclassified():
    cases = [
        ("택배 상자 모서리가 깨져 왔어요", [LABELS[0]]),
        ("완충재 없이 포장되어 파손됐어요", [LABELS[0]]),
        ("배송 박스가 구겨져 도착했어요", [LABELS[0]]),
        ("제품에 금이 가서 교환했어요", [LABELS[0]]),
        ("에어캡이 없어 깨졌습니다", [LABELS[0]]),
        ("택배 중 부딪혀 찌그러졌어요", [LABELS[0]]),
        ("포장 상자가 찢어졌어요", [LABELS[0]]),
        ("뚜껑 파손, 상자도 구겨져 왔어요", [LABELS[0]]),
        ("플라스틱 냄새가 심해요", [LABELS[1]]),
        ("고무 냄새가 오래 남아요", [LABELS[1]]),
        ("입구 마감이 날카롭습니다", [LABELS[1]]),
        ("표면에 스크래치가 있어요", [LABELS[1]]),
        ("코팅이 긁혀서 벗겨졌어요", [LABELS[1]]),
        ("냄새와 거친 마감이 아쉬워요", [LABELS[1]]),
        ("손잡이 칠이 벗겨졌어요", [LABELS[1]]),
        ("고무 악취 때문에 사용하기 힘들어요", [LABELS[1]]),
        ("가방 안에서 음료가 샜어요", [LABELS[2]]),
        ("패킹에서 누수가 생겨요", [LABELS[2]]),
        ("밀폐가 안 돼 물이 흘러요", [LABELS[2]]),
        ("뚜껑을 닫아도 새서 젖었어요", [LABELS[2]]),
        ("실리콘 패킹이 빠져 샘이 있어요", [LABELS[2]]),
        ("누수 때문에 교환을 요청합니다", [LABELS[2]]),
        ("배송 파손과 냄새가 모두 있어요", [LABELS[0], LABELS[1]]),
        ("포장이 깨졌고 가방에서 물도 샜어요", [LABELS[0], LABELS[2]]),
        ("냄새가 나고 패킹 누수도 있어요", [LABELS[1], LABELS[2]]),
        ("박스 파손, 냄새, 누수까지 있어요", [LABELS[0], LABELS[1], LABELS[2]]),
        ("이 제품은 안 새요", ["미분류"]),
        ("물이 새지 않아요", ["미분류"]),
        ("누수가 없고 밀폐 잘 됩니다", ["미분류"]),
        ("색상이 화면과 달라요", ["미분류"]),
    ]
    result = app.analyze(app.normalize([{"text": text} for text, _ in cases]))
    assert len(cases) == 30
    actual = {row["text"]: row["topics"] for group in result["grouped"].values() for row in group}
    assert actual == {text: expected for text, expected in cases}
    assert result["counts"]["미분류"] == 4
    assert result["counts"][LABELS[0]] == 11
    assert result["counts"][LABELS[1]] == 11
    assert result["counts"][LABELS[2]] == 9


def test_original_50_virtual_reviews_keep_representative_counts():
    rows = app.normalize(app.scenario_reviews())
    result = app.analyze(rows)
    assert len(rows) == 50
    assert result["representative_counts"] == {LABELS[0]: 22, LABELS[1]: 15, LABELS[2]: 9, "미분류": 4}
    assert result["percentages"][LABELS[0]] == round(result["counts"][LABELS[0]] / 50 * 100, 1)
    assert len(app.copy_drafts(result)) == 3


@pytest.mark.parametrize("count", [4999, 5000, 5001])
def test_review_limit(count):
    rows = [{"text": "배송 파손 리뷰", "rating": 2}] * count
    if count > app.MAX_REVIEWS:
        with pytest.raises(ValueError):
            app.normalize(rows)
    else:
        assert len(app.normalize(rows)) == count


@pytest.mark.parametrize("length", [3999, 4000, 4001])
def test_review_text_limit(length):
    rows = [{"text": "가" * length}]
    if length > 4000:
        with pytest.raises(ValueError):
            app.normalize(rows)
    else:
        assert len(app.normalize(rows)[0]["text"]) == length


@pytest.mark.parametrize("size", [1999999, 2000000, 2000001])
def test_file_size_limit(size):
    data = b"\n" * size
    if size > app.MAX_FILE_BYTES:
        with pytest.raises(ValueError):
            app.parse_file(data, "reviews.txt")
    else:
        with pytest.raises(ValueError, match="분석할 리뷰가 없습니다"):
            app.parse_file(data, "reviews.txt")


def test_formats_aliases_manual_mapping_and_encodings():
    csv_data = "평점,상품평,작성일,상품명,옵션\n2,배송 중 박스 파손,2026-10-01,텀블러,블랙".encode("utf-8-sig")
    row = app.parse_file(csv_data, "reviews.csv")[0]
    assert (row["rating"], row["text"], row["date"], row["product"], row["option"]) == (2, "배송 중 박스 파손", "2026-10-01", "텀블러", "블랙")
    assert app.parse_file("review\n깨졌어요".encode("utf-8-sig"), "a.csv")[0]["text"] == "깨졌어요"
    cp949 = "review\n포장이 깨졌어요".encode("cp949")
    assert app.parse_file(cp949, "a.csv")[0]["text"] == "포장이 깨졌어요"
    assert app.normalize([{"body": "배송 상자 파손", "score": "1"}], {"text": "body", "rating": "score"})[0]["rating"] == 1
    assert app.normalize([{"text": "한 줄 리뷰"}])[0]["rating"] == 1
    assert app.parse_file(json.dumps({"reviews": [{"review": "JSON 리뷰", "rating": 2}]}, ensure_ascii=False).encode(), "a.json")[0]["text"] == "JSON 리뷰"
    assert app.parse_file(b"TXT review", "a.txt")[0]["text"] == "TXT review"
    book = Workbook(); sheet = book.active
    sheet.append(["별점", "리뷰", "작성일"]); sheet.append([2, "XLSX 리뷰", "2026-10-02"])
    data = io.BytesIO(); book.save(data)
    xlsx = app.parse_file(data.getvalue(), "a.xlsx")[0]
    assert (xlsx["rating"], xlsx["text"], xlsx["date"]) == (2, "XLSX 리뷰", "2026-10-02")


def test_reclassification_updates_counts_and_revert_restores():
    rows = app.normalize([{"text": "상품 색상이 다릅니다"}, {"text": "배송 박스 파손"}])
    baseline = app.analyze(rows)
    overrides = {1: LABELS[0]}
    edited = app.analyze(rows, overrides=overrides)
    assert baseline["counts"]["미분류"] == 1
    assert edited["counts"][LABELS[0]] == baseline["counts"][LABELS[0]] + 1
    assert rows[0]["edited"] is True
    overrides.pop(1)
    reverted = app.analyze(rows, overrides=overrides)
    assert reverted["counts"] == baseline["counts"]


def test_category_dictionary_and_date_product_option_filters():
    rows=app.normalize([
        {"text":"특수한 재질 불만", "rating":1, "작성일":"2026-10-01", "상품명":"컵A", "옵션":"검정"},
        {"text":"배송 박스 파손", "rating":1, "작성일":"2026-10-03", "상품명":"컵B", "옵션":"흰색"},
    ])
    edited_topics={**app.TOPICS,"재질": ["재질"]}
    result=app.analyze(rows,topics=edited_topics,filters={"start":"2026-10-01","end":"2026-10-02","product":"컵A","option":"검정"})
    assert result["selected"] == 1
    assert result["counts"]["재질"] == 1
    assert result["grouped"]["재질"][0]["text"] == "특수한 재질 불만"


def test_export_workbook_values_match_current_analysis():
    rows = app.normalize([{"text": "배송 파손"}, {"text": "제품 냄새"}, {"text": "상품 색상"}])
    overrides = {3: LABELS[1]}
    result = app.analyze(rows, overrides=overrides)
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["유형", "건수", "비율"])
    for topic, count in result["counts"].items():
        sheet.append([topic, count, result["percentages"][topic]])
    values = list(sheet.values)
    observed = {row[0]: (row[1], row[2]) for row in values[1:]}
    assert observed == {topic: (count, result["percentages"][topic]) for topic, count in result["counts"].items()}


def test_streamlit_app_initial_and_analysis_flow():
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file("app.py").run(timeout=30)
    assert not at.exception
    assert next(r for r in at.radio if r.label == "리뷰 입력 방식").options == ["내장 리뷰 50건", "리뷰 파일", "본문 붙여넣기"]
    next(b for b in at.button if b.label == "리뷰 분석 시작").click().run(timeout=30)
    assert not at.exception
    assert at.session_state["result"]["selected"] == 50
