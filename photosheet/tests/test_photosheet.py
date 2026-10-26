import io
import sys
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import (MAX_BYTES, PhotoError, filename_fields, make_excel, paginate,
                 read_photo, sort_photos, dimensions, validate_limits, auto_category,
                 edit_photo, preview_html, slot_geometry, fitted, PAGE_H, HEADER_H, CAPTION_H)


def image_bytes(orientation=1, taken='2026:10:05 10:00:00', fmt='JPEG', size=(120, 80)):
    image = Image.new('RGB', size, 'white')
    # 회전 방향까지 실패할 수 있게 비대칭의 실제 픽셀을 넣습니다.
    image.paste('red', (0, 0, size[0]//2, size[1]//2))
    image.paste('blue', (size[0]//2, size[1]//2, size[0], size[1]))
    exif = Image.Exif()
    exif[274] = orientation
    if taken:
        exif[36867] = taken
    out = io.BytesIO()
    image.save(out, fmt, exif=exif, quality=100)
    return out.getvalue()


def photo(**changes):
    return replace(read_photo('현장_점검.jpg', image_bytes(), 1), **changes)


@pytest.mark.parametrize('orientation,transpose', [(6, Image.Transpose.ROTATE_270), (3, Image.Transpose.ROTATE_180), (8, Image.Transpose.ROTATE_90)])
def test_rotation(orientation, transpose):
    raw = image_bytes(orientation)
    actual = read_photo('회전.jpg', raw, 1)
    expected = Image.open(io.BytesIO(raw)).transpose(transpose)
    result = Image.open(io.BytesIO(actual.jpeg))
    assert result.size == expected.size
    assert actual.rotated
    for x, y in [(10, 10), (result.width-11, result.height-11)]:
        assert sum(abs(a-b) for a, b in zip(result.getpixel((x,y)), expected.getpixel((x,y)))) < 35
    assert not result.getexif().get(274)


def test_date_and_fallback():
    p = photo()
    assert p.taken == '2026-10-05'
    assert p.sort_time == datetime(2026,10,5,10)
    for missing in (None, 'invalid'):
        q = read_photo('none.jpg', image_bytes(taken=missing), datetime(2026,10,20,12).timestamp())
        assert q.taken == '촬영일 없음'
        assert q.sort_time == datetime(2026,10,20,12)


def test_sorting_and_names():
    assert filename_fields('바닥_방수전.jpg') == ('바닥','방수전')
    assert filename_fields('벽_중간_점검.png') == ('벽','중간_점검')
    assert filename_fields('IMG0412.jpg') == ('','')
    assert filename_fields('IMG_0412.jpg') == ('','')
    result = sort_photos([photo(name='c', sort_time=datetime(2026,10,20)), photo(name='b'), photo(name='a')])
    assert [p.name for p in result] == ['a','b','c']
    assert [p.order for p in result] == [1,2,3]


@pytest.mark.parametrize('slots,shape,pages', [(2,(2,1),4),(4,(2,2),2),(6,(3,2),2)])
def test_layout(slots,shape,pages):
    photos = [photo(order=i) for i in range(7)]
    assert dimensions(slots) == shape
    result = paginate(photos,slots)
    assert len(result) == pages
    assert all(len(page) == slots for page in result)
    assert sum(p is not None for page in result for p in page) == 7


def test_grouping():
    photos = [photo(category=c,location=l,order=n) for n,(l,c) in enumerate([
        ('바닥','후'),('벽','중'),('바닥','전'),('바닥','중'),('벽','전'),('바닥','전'),('','전'),('바닥','기타')])]
    result = paginate(photos,4,True)
    assert [p.category if p else None for p in result[0][:3]] == ['전','중','후']
    assert [p.category if p else None for p in result[0][3:]] == ['전',None,None]
    assert all(p.location == '바닥' for p in result[0] if p)
    assert sum(p is not None for page in result for p in page) == len(photos)
    assert all(len(page) == 6 for page in result)


@pytest.mark.parametrize('slots',[2,4,6])
def test_excel(slots):
    photos = [photo(order=i,content='=1+1') for i in range(7)]
    data, pages = make_excel(photos,slots,cover={'title':'가상 공사','date':'2026-10-26','author':''})
    sheet = load_workbook(io.BytesIO(data)).active
    assert len(sheet._images) == 7
    assert len(sheet.row_breaks.brk)+1 == pages
    assert sheet.page_setup.paperSize == int(sheet.PAPERSIZE_A4)
    assert sheet.page_setup.orientation == 'portrait'
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.page_setup.fitToHeight == 0
    assert sheet.sheet_properties.pageSetUpPr.fitToPage
    assert sheet.print_area and sheet.page_margins.left == .3
    assert sheet.oddHeader.center.text == '가상 공사'
    assert sheet.oddHeader.right.text == '&P / &N'
    assert all(c.data_type != 'f' for row in sheet for c in row)
    with ZipFile(io.BytesIO(data)) as archive:
        assert len([n for n in archive.namelist() if n.startswith('xl/media/')]) == 7
    for image in sheet._images:
        assert image.anchor._from.colOff > 0
        assert image.anchor._from.rowOff > 0


def test_reduction():
    image = Image.effect_noise((2600,1800),80).convert('RGB')
    out = io.BytesIO(); image.save(out,'PNG')
    p = read_photo('noise.png',out.getvalue(),1)
    assert max(p.width,p.height) == 1600
    assert len(p.jpeg) < len(out.getvalue())
    excel,_ = make_excel([p])
    assert len(excel) < p.original_size


@pytest.mark.parametrize('count',[199,200,201])
def test_count_boundary(count):
    raw = image_bytes()
    photos = [photo(order=i) for i in range(count)]
    if count > 200:
        with pytest.raises(PhotoError,match='200장 이하'):
            make_excel(photos)
    else:
        validate_limits(count,[len(raw)]*count)
        excel,_ = make_excel(photos)
        assert len(load_workbook(io.BytesIO(excel)).active._images) == count


@pytest.mark.parametrize('size',[MAX_BYTES-1,MAX_BYTES,MAX_BYTES+1])
def test_size_boundary(size):
    raw = image_bytes()
    data = raw + bytes(size-len(raw))
    if size > MAX_BYTES:
        with pytest.raises(PhotoError,match='15MB 以下|15MB 이하'):
            read_photo('boundary.jpg',data,1)
    else:
        p = read_photo('boundary.jpg',data,1)
        assert p.original_size == size


@pytest.mark.parametrize('name,data',[('bad.heic',b'bad'),('bad.jpg',b'broken'),('bad.gif',image_bytes(fmt='PNG')),('bad.jpg',image_bytes(fmt='BMP'))], ids=['heic','broken','suffix','magic'])
def test_errors(name,data):
    with pytest.raises(PhotoError,match='주세요'):
        read_photo(name,data,1)


def test_empty_and_transparency():
    with pytest.raises(PhotoError,match='한 장 이상'):
        validate_limits(0,[])
    with pytest.raises(PhotoError,match='한 장 이상'):
        make_excel([])
    out=io.BytesIO();Image.new('RGBA',(20,20),(255,0,0,0)).save(out,'PNG')
    result=Image.open(io.BytesIO(read_photo('clear.png',out.getvalue(),1).jpeg))
    assert result.getpixel((10,10)) == (255,255,255)


@pytest.mark.parametrize('name,content,expected', [
    ('방수전.jpg', '', '전'), ('전.jpg', '', '전'), ('현장_전.png', '', '전'),
    ('현장(후).jpg', '', '후'), ('시공중.jpg', '', '중'), ('완료후.png', '', '후'),
    ('전기.jpg', '', '기타'), ('후면.jpg', '', '기타'), ('중간.jpg', '', '기타'),
    ('현장_전기_후면.jpg', '', '기타'), ('IMG_0001.jpg', '', '기타'),
    ('IMG_0001.jpg', '방수전', '전'), ('IMG_0001.jpg', '시공중', '중'),
    ('IMG_0001.jpg', '완료후', '후'), ('IMG_0001.jpg', '전기 후면', '기타'),
    ('현장_전_후.jpg', '', '기타'), ('IMG_0001.jpg', '전·중·후', '기타'),
])
def test_category_boundaries(name, content, expected):
    assert auto_category(name, content) == expected
    p = read_photo(name, image_bytes(), 1)
    if content:
        p = edit_photo(p, {'내용': content})
    assert p.category == expected
    excel, _ = make_excel([p])
    assert load_workbook(io.BytesIO(excel)).active['A5'].value.startswith(expected + ' · ')
    assert expected + ' · ' in preview_html([p], 4, False, {})


def test_category_edit_and_manual_choice():
    p = edit_photo(read_photo('IMG_0001.jpg', image_bytes(), 1), {'내용': '방수전'})
    assert p.category == '전'
    p = edit_photo(p, {'내용': '완료후'})
    assert p.category == '후'
    assert edit_photo(p, {'내용': '전기'}).category == '기타'
    p = edit_photo(p, {'구분': '중', '내용': '완료후 확인'})
    assert p.category == '중'
    assert edit_photo(p, {'위치': '벽'}).category == '중'


@pytest.mark.parametrize('slots,grouped', [(2,False),(4,False),(6,False),(4,True)])
def test_a4_photo_geometry(slots, grouped):
    cols, rows = dimensions(slots, grouped)
    slot_w, image_h, slot_h = slot_geometry(slots, grouped)
    assert HEADER_H + rows * slot_h <= PAGE_H
    photos = [read_photo(f'현장{i}_전.jpg', image_bytes(size=size), 1)
              for i, size in enumerate([(400,300), (300,400), (400,299), (400,301), (1600,900), (900,1600)])]
    excel, _ = make_excel(photos, slots, grouped)
    sheet = load_workbook(io.BytesIO(excel)).active
    per_page = 3 + rows * 2
    for page_index, base in enumerate(range(0, sheet.max_row, per_page)):
        _, page_image_h, page_slot_h = slot_geometry(slots, grouped, paginate(photos,slots,grouped)[page_index])
        assert HEADER_H + rows * page_slot_h <= PAGE_H
        heights = [sheet.row_dimensions[r].height for r in range(base+1, base+per_page+1)]
        # A4 세로의 실제 인쇄 가능 높이(0.3인치 여백), 전체 폭 720px로 축소 인쇄.
        scale = (210 / 25.4 * 72 - 2 * .3 * 72) / 720
        assert sum(heights) / .75 * scale <= 297 / 25.4 * 72 - 2 * .3 * 72
        assert sheet.row_dimensions[base+4].height == pytest.approx(page_image_h * .75)
        assert sheet.row_dimensions[base+5].height == CAPTION_H * .75
    for i, drawing in enumerate(sheet._images):
        _, page_image_h, _ = slot_geometry(slots, grouped, paginate(photos,slots,grouped)[i//(cols*rows)])
        anchor = drawing.anchor
        assert anchor._from.colOff >= 4 * 9525 - 1
        assert anchor._from.rowOff >= 4 * 9525 - 1
        assert anchor._from.colOff + anchor.ext.cx <= (slot_w - 4) * 9525 + 1
        assert anchor._from.rowOff + anchor.ext.cy <= (page_image_h - 4) * 9525 + 1
    landscape = photos[0]
    width, height = fitted(landscape, slot_w, image_h)
    assert width == pytest.approx(slot_w - 8)
    assert height == pytest.approx(image_h - 8)
    portrait = photos[1]
    width, height = fitted(portrait, slot_w, image_h)
    assert width / height == pytest.approx(portrait.width / portrait.height)
    html = preview_html(photos, slots, grouped, {})
    _, image_h, slot_h = slot_geometry(slots,grouped,paginate(photos,slots,grouped)[0])
    assert f'grid-template-rows:repeat({rows},{slot_h}px)' in html
    assert f'height:{image_h}px' in html


@pytest.mark.parametrize('size', [(399,300),(400,300),(401,300),(450,300),(300,300)])
@pytest.mark.parametrize('slots', [2,4,6])
def test_landscape_height_tracks_actual_ratio(slots, size):
    photos = [read_photo(f'ratio{i}.jpg', image_bytes(size=size), 1) for i in range(slots)]
    slot_w, image_h, slot_h = slot_geometry(slots,False,photos)
    assert HEADER_H + dimensions(slots)[1] * slot_h <= PAGE_H
    width, height = fitted(photos[0],slot_w,image_h)
    assert width == pytest.approx(slot_w-8)
    assert height == pytest.approx(image_h-8)
